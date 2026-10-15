"""报表接口：出率日报、成本日报、成本链路、异常清单、能耗日报/月报、趋势。全部可导出 Excel。"""
from __future__ import annotations

import calendar
import io
from datetime import date, timedelta
from decimal import Decimal
from types import SimpleNamespace
from urllib.parse import quote

from fastapi import APIRouter, Body, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..calc.engine import Thresholds
from ..db import get_db
from ..models import (
    ENERGIES, STAGES, AnomalyNote, CalcEnergy, CalcIssue, CalcMaterial, CalcProduct, Material, Meter, PackSpec,
    Product,
)
from .common import parse_date

router = APIRouter(prefix="/api/reports")

SOURCE = {"actual": "实际", "std": "标准", "fixed": "固定100%", "missing": "缺数"}


def s(v):
    return None if v is None else str(v)


def rng(start: str | None, end: str | None):
    e = parse_date(end) if end else date.today()
    st = parse_date(start) if start else e
    if st > e:
        st, e = e, st
    if (e - st).days > 400:
        raise HTTPException(400, "日期区间请不要超过 400 天")
    return st, e


def dev_pp(a, b):
    if a is None or b is None:
        return None
    return (a - b) * 100


def alert_keys(db, st, e):
    return {(i.date, i.item_key) for i in db.scalars(
        select(CalcIssue).where(CalcIssue.date >= st, CalcIssue.date <= e, CalcIssue.kind == "alert"))}


# ---------------------------------------------------------------- 出率日报

def yield_report(db: Session, st: date, e: date, all_materials: bool = False):
    th = Thresholds(db)
    mats = {m.id: m for m in db.scalars(select(Material))}
    prods = {p.id: p for p in db.scalars(select(Product))}
    alerts = alert_keys(db, st, e)
    rows = []
    for r in db.scalars(select(CalcMaterial).where(CalcMaterial.date >= st, CalcMaterial.date <= e)
                        .order_by(CalcMaterial.date, CalcMaterial.material_id)):
        m = mats[r.material_id]
        # 类别按计算当天的版本（存在计算结果里），资料后来改了也不影响历史报表
        m = SimpleNamespace(id=m.id, name=m.name, category=r.category or m.category)
        if r.yield_source == "fixed" and not all_materials:
            continue
        key = f"m:{m.id}:process"
        rows.append({"date": r.date.isoformat(), "stage": "process", "stage_name": STAGES["process"],
                     "name": m.name, "group": m.category, "input": s(r.qty1), "output": s(r.qty2),
                     "actual": s(r.actual_yield), "std": s(r.std_yield),
                     "dev": s(dev_pp(r.actual_yield, r.std_yield)) if r.yield_source == "actual" else None,
                     "source": SOURCE.get(r.yield_source, r.yield_source), "threshold": s(th.for_material(m)),
                     "alert": (r.date, key) in alerts, "key": key})
    for r in db.scalars(select(CalcProduct).where(CalcProduct.date >= st, CalcProduct.date <= e)
                        .order_by(CalcProduct.date, CalcProduct.product_id)):
        p = prods[r.product_id]
        p = SimpleNamespace(id=p.id, name=p.name, ptype=r.ptype or p.ptype, team=r.team or p.team,
                            cooled=p.cooled if r.cooled is None else r.cooled)
        stage = "oil" if p.ptype == "麻油" else "cook"
        key = f"p:{p.id}:{stage}"
        rows.append({"date": r.date.isoformat(), "stage": stage, "stage_name": STAGES[stage], "name": p.name,
                     "group": p.team, "input": s(r.input_weight), "output": s(r.output_weight),
                     "actual": s(r.cook_yield_actual), "std": s(r.cook_yield_std),
                     "dev": s(dev_pp(r.cook_yield_actual, r.cook_yield_std)), "source": "实际",
                     "threshold": s(th.for_product(p.id, stage)), "alert": (r.date, key) in alerts, "key": key})
        if p.cooled and r.cool_source:
            key = f"p:{p.id}:cool"
            rows.append({"date": r.date.isoformat(), "stage": "cool", "stage_name": STAGES["cool"], "name": p.name,
                         "group": p.team, "input": s(r.output_weight), "output": s(r.cooled_weight),
                         "actual": s(r.cool_yield_actual), "std": s(r.cool_yield_std),
                         "dev": s(dev_pp(r.cool_yield_actual, r.cool_yield_std)) if r.cool_source == "actual" else None,
                         "source": SOURCE.get(r.cool_source, r.cool_source),
                         "threshold": s(th.for_product(p.id, "cool")), "alert": (r.date, key) in alerts, "key": key})
        if r.packed_weight is not None:
            key = f"p:{p.id}:pack"
            rows.append({"date": r.date.isoformat(), "stage": "pack", "stage_name": STAGES["pack"], "name": p.name,
                         "group": p.team, "input": s(r.prepack_weight), "output": s(r.packed_weight),
                         "actual": s(r.pack_yield_actual), "std": s(r.pack_yield_std),
                         "dev": s(dev_pp(r.pack_yield_actual, r.pack_yield_std)), "source": "实际",
                         "threshold": s(th.for_product(p.id, "pack")), "alert": (r.date, key) in alerts, "key": key})
    cols = [("date", "日期", "text"), ("stage_name", "环节", "text"), ("name", "原料/产品", "text"),
            ("group", "类别/班组", "text"), ("input", "投入(斤)", "weight"), ("output", "产出(斤)", "weight"),
            ("actual", "实际出率", "yield"), ("std", "标准出率", "yield"), ("dev", "偏差(百分点)", "pp"),
            ("threshold", "阈值(百分点)", "pp"), ("source", "取值", "text"), ("alert", "超限", "bool")]
    return {"title": "出率日报", "columns": cols, "rows": rows}


@router.get("/yield")
def get_yield(start: str | None = None, end: str | None = None, all_materials: bool = False,
              db: Session = Depends(get_db)):
    return yield_report(db, *rng(start, end), all_materials)


# ---------------------------------------------------------------- 成本日报

def cost_report(db: Session, st: date, e: date, ptype: str | None = None, kw: str | None = None):
    """ptype / kw：和页面上的筛选一致，导出时带上，导出的就是页面上看到的。"""
    prods = {p.id: p for p in db.scalars(select(Product))}
    rows = []
    for r in db.scalars(select(CalcProduct).where(CalcProduct.date >= st, CalcProduct.date <= e)
                        .order_by(CalcProduct.date, CalcProduct.product_id)):
        p = prods[r.product_id]
        if (ptype and (r.ptype or p.ptype) != ptype) or (kw and kw not in p.name):
            continue
        rows.append({"date": r.date.isoformat(), "product_id": p.id, "name": p.name, "ptype": r.ptype or p.ptype,
                     "team": r.team or p.team, "ok": r.ok, "input_cost": s(r.input_cost_actual),
                     "out_a": s(r.out_price_actual), "out_s": s(r.out_price_std),
                     "cooled_a": s(r.cooled_price_actual), "cooled_s": s(r.cooled_price_std),
                     "final_a": s(r.final_price_actual), "final_s": s(r.final_price_std),
                     "bag_fee": s(r.bag_fee), "weight": s(r.packed_weight or r.cooled_weight or r.output_weight),
                     "diff_price": s(r.diff_price), "diff_amount": s(r.diff_amount),
                     "loss_amount": s(r.loss_amount),
                     "quality": (r.quality or {}).get("label") or ("算出" if r.ok else "算不出"),
                     "quality_status": (r.quality or {}).get("status"),
                     "quality_flags": "；".join((r.quality or {}).get("flags", []))})
    cols = [("date", "日期", "text"), ("name", "产品", "text"), ("ptype", "类型", "text"), ("team", "班组", "text"),
            ("input_cost", "投料成本(元)", "money"), ("out_a", "出机单价·实际", "price"), ("out_s", "出机单价·标准", "price"),
            ("cooled_a", "冷后单价·实际", "price"), ("cooled_s", "冷后单价·标准", "price"),
            ("final_a", "成品单价·实际", "price"), ("final_s", "成品单价·标准", "price"),
            ("bag_fee", "袋费(元)", "money"), ("weight", "计价重量(斤)", "weight"),
            ("diff_price", "差额(元/斤)", "price"), ("diff_amount", "当日差额(元)", "money"),
            ("quality", "数据状态", "text"), ("loss_amount", "全损损失(元)", "money"), ("quality_flags", "状态说明", "text")]
    return {"title": "成本日报", "columns": cols, "rows": rows}


@router.get("/cost")
def get_cost(start: str | None = None, end: str | None = None, ptype: str | None = None, kw: str | None = None,
             db: Session = Depends(get_db)):
    return cost_report(db, *rng(start, end), ptype, kw)


# ---------------------------------------------------------------- 成本链路

@router.get("/trace")
def trace(product_id: int, day: str, db: Session = Depends(get_db)):
    d = parse_date(day)
    r = db.scalars(select(CalcProduct).where(CalcProduct.product_id == product_id, CalcProduct.date == d)).first()
    if r is None:
        raise HTTPException(404, "这一天没有这个产品的计算结果")
    p = db.get(Product, product_id)
    specs = {sp.id: sp.name for sp in db.scalars(select(PackSpec))}
    t = dict(r.trace or {})
    for b in t.get("bags", []):
        b["spec"] = specs.get(b["spec_id"])
    issues = [{"kind": i.kind, "message": i.message} for i in db.scalars(
        select(CalcIssue).where(CalcIssue.date == d, CalcIssue.item_key.like(f"p:{product_id}:%")))]
    calc = {c: s(getattr(r, c)) for c in (
        "input_weight", "input_cost_actual", "input_cost_std", "output_weight", "cook_yield_actual", "cook_yield_std",
        "out_price_actual", "out_price_std", "cooled_weight", "cool_yield_actual", "cool_yield_std", "cool_source",
        "cooled_price_actual", "cooled_price_std", "prepack_weight", "packed_weight", "pack_yield_actual",
        "pack_yield_std", "bag_fee", "final_price_actual", "final_price_std", "diff_price", "diff_amount")}
    snap = {"ptype": r.ptype or p.ptype, "team": r.team or p.team,
            "cooled": p.cooled if r.cooled is None else r.cooled, "packed": p.packed if r.packed is None else r.packed}
    return {"date": d.isoformat(), "product": {"id": p.id, "name": p.name, **snap},
            "ok": r.ok, "calc": calc, "trace": t, "attribution": r.attribution or [], "issues": issues,
            "quality": r.quality, "loss_amount": s(r.loss_amount)}


# ---------------------------------------------------------------- 异常清单

def issue_report(db: Session, st: date, e: date, kind: str | None = None):
    stmt = select(CalcIssue).where(CalcIssue.date >= st, CalcIssue.date <= e)
    if kind:
        stmt = stmt.where(CalcIssue.kind == kind)
    notes = {(n.date, n.item_key): n.note for n in db.scalars(
        select(AnomalyNote).where(AnomalyNote.date >= st, AnomalyNote.date <= e))}
    rows = [{"id": i.id, "date": i.date.isoformat(), "kind": i.kind, "kind_name": "超限" if i.kind == "alert" else "提示",
             "item_key": i.item_key, "name": i.object_name, "stage": i.stage,
             "stage_name": STAGES.get(i.stage, "") if i.stage else "", "actual": s(i.actual), "std": s(i.std),
             "dev": s(i.dev_pp), "threshold": s(i.threshold), "message": i.message,
             "note": notes.get((i.date, i.item_key))}
            for i in db.scalars(stmt.order_by(CalcIssue.date, CalcIssue.kind, CalcIssue.id))]
    cols = [("date", "日期", "text"), ("kind_name", "类型", "text"), ("name", "对象", "text"),
            ("stage_name", "环节", "text"), ("actual", "实际", "yield"), ("std", "标准", "yield"),
            ("dev", "偏差(百分点)", "pp"), ("message", "说明", "text"), ("note", "原因备注", "text")]
    return {"title": "异常清单", "columns": cols, "rows": rows}


@router.get("/issues")
def get_issues(start: str | None = None, end: str | None = None, kind: str | None = None,
               db: Session = Depends(get_db)):
    return issue_report(db, *rng(start, end), kind)


@router.put("/notes")
def put_note(payload: dict = Body(...), db: Session = Depends(get_db)):
    d = parse_date(payload["date"])
    key = payload["item_key"]
    text = (payload.get("note") or "").strip()
    row = db.scalars(select(AnomalyNote).where(AnomalyNote.date == d, AnomalyNote.item_key == key)).first()
    if not text:
        if row:
            db.delete(row)
    elif row is None:
        db.add(AnomalyNote(date=d, item_key=key, note=text))
    else:
        row.note = text
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- 能耗

def energy_rows(db: Session, st: date, e: date):
    return list(db.scalars(select(CalcEnergy).where(CalcEnergy.date >= st, CalcEnergy.date <= e)
                           .order_by(CalcEnergy.date)))


def dec(v):
    return Decimal(v) if v not in (None, "") else None


def energy_daily(db: Session, st: date, e: date):
    rows = []
    for r in energy_rows(db, st, e):
        sm = r.data["summary"]
        rows.append({
            "date": r.date.isoformat(), "output_jin": s(r.total_output), "output_ton": r.data.get("output_ton"),
            "gas": sm["gas"]["plant"], "elec": sm["elec"]["plant"], "water": sm["water"]["plant"],
            "water_ws": sm["water"]["workshop"],
            "gas_t": sm["gas"]["per_ton"], "elec_t": sm["elec"]["per_ton"], "water_t": sm["water"]["per_ton"],
            "water_ws_t": sm["water"]["workshop_per_ton"], "water_out_t": sm["water"]["outside_per_ton"],
            "gas_cost": sm["gas"]["cost"], "elec_cost": sm["elec"]["cost"], "water_cost": sm["water"]["cost"],
            "cost": r.data.get("cost_total"),
            "complete": r.data.get("complete", all(v["complete"] for v in sm.values()) and r.data.get("has_readings", False)),
            "notes": "；".join(r.data.get("notes", [])),
        })
    cols = [("date", "日期", "text"), ("output_jin", "总产量(斤)", "weight"), ("output_ton", "总产量(吨)", "ton"),
            ("gas", "用气(Nm³)", "qty"), ("elec", "用电(度)", "qty"), ("water", "全厂用水(吨)", "qty"),
            ("water_ws", "车间用水(吨)", "qty"), ("gas_t", "吨耗气(Nm³/吨)", "qty"), ("elec_t", "吨耗电(度/吨)", "qty"),
            ("water_t", "吨耗水", "qty"), ("water_ws_t", "车间吨耗水", "qty"), ("water_out_t", "车间外吨耗水", "qty"),
            ("gas_cost", "燃气费(元)", "money"), ("elec_cost", "电费(元)", "money"), ("water_cost", "水费(元)", "money"),
            ("cost", "能耗费用合计(元)", "money"), ("notes", "说明（不完整原因）", "text")]
    return {"title": "能耗日报", "columns": cols, "rows": rows}


@router.get("/energy-daily")
def get_energy_daily(start: str | None = None, end: str | None = None, db: Session = Depends(get_db)):
    return energy_daily(db, *rng(start, end))


@router.get("/energy-day")
def get_energy_day(day: str, db: Session = Depends(get_db)):
    """某一天的抄表明细。"""
    r = db.get(CalcEnergy, parse_date(day))
    if r is None:
        raise HTTPException(404, "这一天没有能耗数据")
    return {"date": r.date.isoformat(), "total_output": s(r.total_output), **r.data}


def energy_month(db: Session, month: str):
    try:
        y, m = (int(x) for x in month.split("-"))
        st = date(y, m, 1)
    except Exception:
        raise HTTPException(400, "月份格式应为 YYYY-MM")
    e = date(y, m, calendar.monthrange(y, m)[1])
    days = [st + timedelta(days=i) for i in range((e - st).days + 1)]
    recs = {r.date: r for r in energy_rows(db, st, e)}
    meters = list(db.scalars(select(Meter).order_by(Meter.energy, Meter.sort, Meter.id)))
    order = {"gas": 0, "elec": 1, "water": 2}
    meters.sort(key=lambda x: (order.get(x.energy, 9), x.sort, x.id))

    def per_day(fn):
        vals = []
        for d in days:
            r = recs.get(d)
            vals.append(fn(r) if r else None)
        return vals

    def total(vals):
        xs = [dec(v) for v in vals if v is not None]
        return s(sum(xs)) if xs else None

    lines = []
    for mt in meters:
        def usage(r, mid=mt.id):
            for x in r.data["meters"]:
                if x["meter_id"] == mid:
                    return x["usage"]
            return None
        vals = per_day(usage)
        if all(v is None for v in vals):
            continue
        info = next((x for r in recs.values() for x in r.data["meters"] if x["meter_id"] == mt.id), {})
        mode = {"+": "计入", "-": "扣减", "ref": "参考", "": ""}
        lines.append({"label": f"{ENERGIES[mt.energy][0]}·{mt.name}", "group": mt.energy, "kind": "meter",
                      "note": "全厂" + mode.get(info.get("plant", ""), "") + (
                          "，车间" + mode[info["workshop"]] if info.get("workshop") else ""),
                      "values": vals, "total": total(vals)})
    out_vals = per_day(lambda r: s(r.total_output))
    out_total = total(out_vals)
    ton_total = dec(out_total) / 2000 if out_total and dec(out_total) else None
    for ek, label, scope in [("gas", "用气合计(Nm³)", "plant"), ("elec", "用电合计(度)", "plant"),
                             ("water", "全厂用水合计(吨)", "plant"), ("water", "车间用水合计(吨)", "workshop")]:
        vals = per_day(lambda r: r.data["summary"][ek][scope])
        lines.append({"label": label, "group": ek, "kind": "total", "values": vals, "total": total(vals),
                      "_key": (ek, scope)})
    lines.append({"label": "当日总产量(斤)", "group": "out", "kind": "total", "values": out_vals, "total": out_total})
    totals = {l["_key"]: dec(l["total"]) for l in lines if "_key" in l}
    # 月度指标 = 月合计用量 ÷ 月合计产量，不是日指标平均
    for ek, label, scope, field in [("gas", "吨耗气(Nm³/吨)", "plant", "per_ton"), ("elec", "吨耗电(度/吨)", "plant", "per_ton"),
                                    ("water", "吨耗水(吨/吨)", "plant", "per_ton"),
                                    ("water", "车间吨耗水(吨/吨)", "workshop", "workshop_per_ton")]:
        vals = per_day(lambda r: r.data["summary"][ek][field])
        t = totals.get((ek, scope))
        lines.append({"label": label, "group": ek, "kind": "index", "values": vals,
                      "total": s(t / ton_total) if t is not None and ton_total else None})
    wp, ww = totals.get(("water", "plant")), totals.get(("water", "workshop"))
    lines.append({"label": "车间外吨耗水(吨/吨)", "group": "water", "kind": "index",
                  "values": per_day(lambda r: r.data["summary"]["water"]["outside_per_ton"]),
                  "total": s((wp - ww) / ton_total) if wp is not None and ww is not None and ton_total else None})
    for ek, label in [("gas", "燃气费(元)"), ("elec", "电费(元)"), ("water", "水费(元)")]:
        vals = per_day(lambda r: r.data["summary"][ek]["cost"])
        lines.append({"label": label, "group": ek, "kind": "cost", "values": vals, "total": total(vals)})
    vals = per_day(lambda r: r.data.get("cost_total"))
    lines.append({"label": "能耗费用合计(元)", "group": "cost", "kind": "cost", "values": vals, "total": total(vals)})
    for l in lines:
        l.pop("_key", None)
    return {"title": f"能耗月报 {month}", "month": month, "days": [d.isoformat() for d in days], "lines": lines,
            "notes": month_coverage(recs, days)}


def month_coverage(recs: dict, days: list[date]) -> list[str]:
    """月报的覆盖范围说明：月合计是“按抄表入账日累计”，漏抄、跨日、跨月、月末没读数都写出来。"""
    notes = ["月合计按抄表入账日累计：某天的读数差记在抄表那天，漏抄的日子不平均摊分。"]
    have = [d for d in days if d in recs and recs[d].data.get("has_readings")]
    if not have:
        return notes + ["本月没有抄表数据。"]
    full = [d for d in have if recs[d].data.get("complete")]
    notes.append(f"有抄表 {len(have)} 天，其中合计完整 {len(full)} 天。")
    no_read = [d for d in days if d in recs and not recs[d].data.get("has_readings") and recs[d].total_output]
    if no_read:
        notes.append("这些天有产量但没抄表：" + "、".join(f"{d:%m-%d}" for d in no_read) + "。")
    seen = set()
    for d in have:
        for mt in recs[d].data["meters"]:
            if mt.get("cross_month") and (d, mt["name"]) not in seen:
                seen.add((d, mt["name"]))
                notes.append(f"{d:%m-%d} {mt['name']}：{mt['note']}，本月合计里含上月的部分。")
    last = have[-1]
    if last < days[-1]:
        notes.append(f"最后一次抄表是 {last:%m-%d}，之后的用量还没入账，本月合计只到这天。")
    return notes


@router.get("/energy-month")
def get_energy_month(month: str, db: Session = Depends(get_db)):
    return energy_month(db, month)


# ---------------------------------------------------------------- 趋势

@router.get("/trend")
def trend(kind: str, start: str | None = None, end: str | None = None, ref_id: int | None = None,
          stage: str | None = None, metric: str | None = None, db: Session = Depends(get_db)):
    """kind: material（原料出率）/ product_yield（产品环节出率）/ product_cost（产品单价）/ energy（能耗指标）。"""
    st, e = rng(start, end)
    pts = []
    if kind == "material":
        for r in db.scalars(select(CalcMaterial).where(CalcMaterial.material_id == ref_id, CalcMaterial.date >= st,
                                                       CalcMaterial.date <= e).order_by(CalcMaterial.date)):
            pts.append({"date": r.date.isoformat(), "actual": s(r.actual_yield) if r.yield_source == "actual" else None,
                        "std": s(r.std_yield)})
        unit = "yield"
    elif kind in ("product_yield", "product_cost"):
        field = {"cook": "cook_yield", "oil": "cook_yield", "cool": "cool_yield", "pack": "pack_yield",
                 "out": "out_price", "cooled": "cooled_price", "final": "final_price"}.get(stage or "", None)
        if field is None:
            raise HTTPException(400, "stage 无效")
        for r in db.scalars(select(CalcProduct).where(CalcProduct.product_id == ref_id, CalcProduct.date >= st,
                                                      CalcProduct.date <= e).order_by(CalcProduct.date)):
            pts.append({"date": r.date.isoformat(), "actual": s(getattr(r, field + "_actual")),
                        "std": s(getattr(r, field + "_std"))})
        unit = "yield" if kind == "product_yield" else "price"
    elif kind == "energy":
        ek, _, f = (metric or "gas.per_ton").partition(".")
        if ek not in ENERGIES:
            raise HTTPException(400, "metric 无效")
        for r in energy_rows(db, st, e):
            pts.append({"date": r.date.isoformat(), "actual": r.data["summary"][ek].get(f or "per_ton"), "std": None})
        unit = "qty"
    else:
        raise HTTPException(400, "kind 无效")
    return {"points": pts, "unit": unit}


# ---------------------------------------------------------------- 导出 Excel

FMT = {"weight": "0.00", "price": "0.0000", "yield": "0.00%", "pp": "0.00", "money": "0.00", "qty": "0.00",
       "ton": "0.0000"}
HEAD_FILL = PatternFill("solid", fgColor="DDE7F0")
RED = Font(color="C00000", bold=True)


def to_cell(v, typ):
    if v is None:
        return None
    if typ == "bool":
        return "是" if v else ""
    if typ in FMT:
        try:
            return float(Decimal(str(v)))
        except Exception:
            return v
    return v


def table_xlsx(title: str, columns, rows, sub: str = ""):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:30]
    ws.append([title + (f"（{sub}）" if sub else "")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([c[1] for c in columns])
    for i in range(1, len(columns) + 1):
        c = ws.cell(2, i)
        c.fill = HEAD_FILL
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = 14 if columns[i - 1][2] != "text" else 18
    for r in rows:
        ws.append([to_cell(r.get(k), t) for k, _, t in columns])
        rr = ws.max_row
        for i, (k, _, t) in enumerate(columns, start=1):
            if t in FMT:
                ws.cell(rr, i).number_format = FMT[t]
        if r.get("alert") or r.get("kind") == "alert":
            for i in range(1, len(columns) + 1):
                ws.cell(rr, i).font = RED
    ws.freeze_panes = "A3"
    return wb


def respond(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}.xlsx"})


@router.get("/export/{name}")
def export(name: str, start: str | None = None, end: str | None = None, month: str | None = None,
           product_id: int | None = None, day: str | None = None, kind: str | None = None,
           all_materials: bool = False, ptype: str | None = None, kw: str | None = None,
           db: Session = Depends(get_db)):
    if name == "energy-month":
        rep = energy_month(db, month or date.today().strftime("%Y-%m"))
        wb = Workbook()
        ws = wb.active
        ws.title = "能耗月报"
        ws.append([rep["title"]])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["项目", "说明"] + [d[8:] + "日" for d in rep["days"]] + ["月合计"])
        for c in ws[2]:
            c.fill, c.font = HEAD_FILL, Font(bold=True)
        for l in rep["lines"]:
            fmt = "0.0000" if l["kind"] == "index" else "0.00"
            ws.append([l["label"], l.get("note", "")] + [to_cell(v, "qty") for v in l["values"]] + [to_cell(l["total"], "qty")])
            for c in ws[ws.max_row][2:]:
                c.number_format = fmt
            if l["kind"] != "meter":
                ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 18
        ws.freeze_panes = "C3"
        return respond(wb, rep["title"])
    if name == "trace":
        if not (product_id and day):
            raise HTTPException(400, "需要产品和日期")
        t = trace(product_id, day, db)
        cols = [("line_no", "行号", "text"), ("name", "投料", "text"), ("kind", "类型", "text"),
                ("qty", "每锅数量(斤)", "weight"), ("price1", "1级单价", "price"), ("yield_actual", "实际出率", "yield"),
                ("yield_std", "标准出率", "yield"), ("yield_source", "出率取值", "text"),
                ("price_actual", "实际单价", "price"), ("price_std", "标准单价", "price"),
                ("cost_actual", "当日投料成本(元)", "money")]
        rows = t["trace"].get("lines", [])
        for r in rows:
            r["yield_source"] = SOURCE.get(r.get("yield_source"), r.get("yield_source"))
        wb = table_xlsx(f"成本链路 {t['product']['name']}", cols, rows, t["date"])
        ws = wb.active
        ws.append([])
        labels = {"input_weight": "投料合计(斤)", "input_cost_actual": "投料成本·实际(元)", "output_weight": "出锅/出机重(斤)",
                  "cook_yield_actual": "熟制出率·实际", "cook_yield_std": "熟制出率·标准", "out_price_actual": "出机单价·实际",
                  "out_price_std": "出机单价·标准", "cooled_weight": "打冷出机重(斤)", "cool_yield_actual": "制冷出率·实际",
                  "cool_yield_std": "制冷出率·标准", "cooled_price_actual": "冷后单价·实际", "cooled_price_std": "冷后单价·标准",
                  "prepack_weight": "包装前重量(斤)", "packed_weight": "包装入库量(斤)", "pack_yield_actual": "包装出率·实际",
                  "pack_yield_std": "包装出率·标准", "bag_fee": "袋费(元)", "final_price_actual": "成品单价·实际",
                  "final_price_std": "成品单价·标准", "diff_price": "差额(元/斤)", "diff_amount": "当日差额(元)"}
        for k, label in labels.items():
            v = t["calc"].get(k)
            if v is not None:
                ws.append([label, to_cell(v, "price")])
                ws.cell(ws.max_row, 2).number_format = "0.00%" if "yield" in k else "0.0000"
        if t["attribution"]:
            ws.append([])
            ws.append(["差额归因", "元/斤", "当日金额(元)"])
            for a in t["attribution"]:
                ws.append([a["factor"], to_cell(a["per_jin"], "price"), to_cell(a.get("amount"), "money")])
        return respond(wb, f"成本链路_{t['product']['name']}_{t['date']}")
    st, e = rng(start, end)
    sub = st.isoformat() if st == e else f"{st} 至 {e}"
    if name == "yield":
        rep = yield_report(db, st, e, all_materials)
    elif name == "cost":
        rep = cost_report(db, st, e, ptype, kw)
    elif name == "issues":
        rep = issue_report(db, st, e, kind)
    elif name == "energy-daily":
        rep = energy_daily(db, st, e)
    else:
        raise HTTPException(404, "没有这个报表")
    return respond(table_xlsx(rep["title"], rep["columns"], rep["rows"], sub), f"{rep['title']}_{sub}")


# ---------------------------------------------------------------- 月度汇总

def month_range(month: str):
    try:
        y, m = (int(x) for x in month.split("-"))
        st = date(y, m, 1)
    except Exception:
        raise HTTPException(400, "月份格式应为 YYYY-MM")
    return st, date(y, m, calendar.monthrange(y, m)[1])


def _ratio(a, b):
    return a / b if a is not None and b else None


def _cost_basis(p: Product, r: CalcProduct):
    """月度汇总按哪一级单价和重量计：与成本差额口径一致（成品看成品，半成品油看冷后，其他看出机）。

    类型、是否制冷/包装按计算当天的设置（结果里存了），没存的旧结果用资料里的当前值。
    """
    ptype = r.ptype or p.ptype
    packed = p.packed if r.packed is None else r.packed
    cooled = p.cooled if r.cooled is None else r.cooled
    if ptype == "成品馅" and packed:
        return r.final_price_actual, r.final_price_std, r.packed_weight
    if ptype != "麻油" and cooled:
        return r.cooled_price_actual, r.cooled_price_std, r.cooled_weight
    return r.out_price_actual, r.out_price_std, r.output_weight


def _product_month(db: Session, st: date, e: date, prods: dict, loss: set = frozenset()):
    out: dict[int, dict] = {}
    for r in db.scalars(select(CalcProduct).where(CalcProduct.date >= st, CalcProduct.date <= e)):
        p = prods[r.product_id]
        a = out.setdefault(p.id, {"days": 0, "cost_days": 0, "weight": Decimal(0), "cost_weight": Decimal(0),
                                  "cost_a": Decimal(0), "cost_s": Decimal(0), "diff_amount": Decimal(0),
                                  "in_w": Decimal(0), "out_w": Decimal(0), "in_w_std": Decimal(0), "out_w_std": Decimal(0),
                                  "loss_days": 0, "loss_cost": Decimal(0), "input_cost": Decimal(0)})
        a["days"] += 1
        pa, ps, w = _cost_basis(p, r)
        if w is not None:
            a["weight"] += w
        if pa is not None and ps is not None and w:
            a["cost_days"] += 1
            a["cost_weight"] += w
            a["cost_a"] += pa * w
            a["cost_s"] += ps * w
        if r.diff_amount is not None:
            a["diff_amount"] += r.diff_amount
        if r.input_cost_actual is not None:
            a["input_cost"] += r.input_cost_actual
        if r.input_weight and r.output_weight is not None:
            a["in_w"] += r.input_weight
            a["out_w"] += r.output_weight
            if r.cook_yield_std is not None:
                a["in_w_std"] += r.input_weight
                a["out_w_std"] += r.input_weight * r.cook_yield_std
        if (p.id, r.date) in loss or (r.quality or {}).get("status") == "loss":
            # 任一环节产出为 0 的整批报废，损失按当天投料成本计
            a["loss_days"] += 1
            if r.input_cost_actual is not None:
                a["loss_cost"] += r.input_cost_actual
    for a in out.values():
        a["price_a"] = _ratio(a["cost_a"], a["cost_weight"])
        a["price_s"] = _ratio(a["cost_s"], a["cost_weight"])
        a["yield_a"] = _ratio(a["out_w"], a["in_w"])
        a["yield_s"] = _ratio(a["out_w_std"], a["in_w_std"])
    return out


def monthly_report(db: Session, month: str):
    """月度汇总：各产品月产量、按产量加权的实际/标准单价和月差额，各原料月出率和出率偏差金额。

    月度单价 = Σ（每天单价 × 当天计价重量）÷ Σ 计价重量，即月总成本 ÷ 月总产量，不是日单价的简单平均；
    月出率 = Σ 产出 ÷ Σ 投入；月标准出率按投入量加权。
    """
    st, e = month_range(month)
    pst, pe = month_range((st - timedelta(days=1)).strftime("%Y-%m"))
    prods = {p.id: p for p in db.scalars(select(Product))}
    mats = {m.id: m for m in db.scalars(select(Material))}
    alerts: dict[str, set] = {}
    loss = set()
    for i in db.scalars(select(CalcIssue).where(CalcIssue.date >= st, CalcIssue.date <= e, CalcIssue.kind == "alert")):
        parts = i.item_key.split(":")
        if len(parts) >= 2:
            alerts.setdefault(f"{parts[0]}:{parts[1]}", set()).add(i.date)
            if parts[0] == "p" and parts[-1] == "loss":
                loss.add((int(parts[1]), i.date))

    cur, prev = _product_month(db, st, e, prods, loss), _product_month(db, pst, pe, prods)
    products = []
    for pid, a in cur.items():
        p = prods[pid]
        pv = prev.get(pid, {}).get("price_a")
        products.append({
            "product_id": pid, "name": p.name, "ptype": p.ptype, "team": p.team,
            "days": a["days"], "no_cost_days": a["days"] - a["cost_days"], "alert_days": len(alerts.get(f"p:{pid}", ())),
            "weight": s(a["weight"]), "input_cost": s(a["input_cost"]),
            "price_a": s(a["price_a"]), "price_s": s(a["price_s"]),
            "diff_price": s(a["price_a"] - a["price_s"]) if a["price_a"] is not None and a["price_s"] is not None else None,
            "diff_amount": s(a["diff_amount"]),
            "prev_price_a": s(pv), "price_change": s(a["price_a"] / pv - 1) if a["price_a"] is not None and pv else None,
            "yield_a": s(a["yield_a"]), "yield_s": s(a["yield_s"]), "yield_dev": s(dev_pp(a["yield_a"], a["yield_s"])),
            "loss_days": a["loss_days"], "loss_cost": s(a["loss_cost"]) if a["loss_days"] else None,
        })
    products.sort(key=lambda x: (-abs(Decimal(x["diff_amount"] or 0)), x["name"]))

    mrows: dict[int, dict] = {}
    for r in db.scalars(select(CalcMaterial).where(CalcMaterial.date >= st, CalcMaterial.date <= e)):
        m = mats.get(r.material_id)
        if m is None:
            continue
        if not (m.needs_process if r.needs_process is None else r.needs_process):  # 按计算当天的设置
            continue
        a = mrows.setdefault(m.id, {"days": 0, "std_days": 0, "q1": Decimal(0), "q2": Decimal(0), "q1_std": Decimal(0),
                                    "q2_std": Decimal(0), "impact": Decimal(0), "impact_days": 0})
        if r.yield_source == "actual" and r.qty1:
            a["days"] += 1
            a["q1"] += r.qty1
            a["q2"] += r.qty2 or 0
            if r.std_yield:
                a["q1_std"] += r.qty1
                a["q2_std"] += r.qty1 * r.std_yield
                if r.price1 is not None:
                    # 同样的 2 级产出，按标准出率只需要 qty2 ÷ 标准出率 的 1 级原料，多用的部分折成金额
                    a["impact"] += r.price1 * (r.qty1 - (r.qty2 or 0) / r.std_yield)
                    a["impact_days"] += 1
        elif r.used and r.yield_source in ("std", "missing"):
            a["std_days"] += 1
    materials = []
    for mid, a in mrows.items():
        m = mats[mid]
        ya, ys = _ratio(a["q2"], a["q1"]), _ratio(a["q2_std"], a["q1_std"])
        materials.append({
            "material_id": mid, "name": m.name, "category": m.category, "days": a["days"],
            "std_days": a["std_days"], "alert_days": len(alerts.get(f"m:{mid}", ())),
            "qty1": s(a["q1"]) if a["days"] else None, "qty2": s(a["q2"]) if a["days"] else None,
            "yield_a": s(ya), "yield_s": s(ys), "yield_dev": s(dev_pp(ya, ys)),
            "impact": s(a["impact"]) if a["impact_days"] else None,
        })
    materials.sort(key=lambda x: (-abs(Decimal(x["impact"] or 0)), x["name"]))

    energy = list(db.scalars(select(CalcEnergy).where(CalcEnergy.date >= st, CalcEnergy.date <= e)))
    out_total = sum((r.total_output or 0 for r in energy), start=Decimal(0))
    tons = out_total / 2000 if out_total else None
    usage = {k: sum((Decimal(r.data["summary"][k]["plant"]) for r in energy
                     if r.data.get("summary", {}).get(k, {}).get("plant") is not None), start=Decimal(0))
             for k in ENERGIES}
    energy_cost = sum((Decimal(r.data["cost_total"]) for r in energy if r.data.get("cost_total") is not None),
                      start=Decimal(0))
    summary = {
        "production_days": sum(1 for r in energy if r.total_output),
        "output": s(out_total), "output_ton": s(tons),
        "diff_amount": s(sum((Decimal(x["diff_amount"]) for x in products if x["diff_amount"] is not None),
                             start=Decimal(0))),
        "material_impact": s(sum((Decimal(x["impact"]) for x in materials if x["impact"] is not None), start=Decimal(0))),
        "loss_cost": s(sum((Decimal(x["loss_cost"]) for x in products if x["loss_cost"] is not None), start=Decimal(0))),
        "energy_cost": s(energy_cost),
        "energy_per_ton": {k: s(usage[k] / tons) if tons else None for k in ENERGIES},
        "energy_cost_per_ton": s(energy_cost / tons) if tons else None,
        "alert_days": len({d for ds in alerts.values() for d in ds}),
    }
    return {"title": f"月度汇总 {month}", "month": month, "start": st.isoformat(), "end": e.isoformat(),
            "prev_month": pst.strftime("%Y-%m"), "summary": summary, "products": products, "materials": materials,
            "product_columns": MONTH_PRODUCT_COLS, "material_columns": MONTH_MATERIAL_COLS}


MONTH_PRODUCT_COLS = [
    ("name", "产品", "text"), ("ptype", "类型", "text"), ("team", "班组", "text"), ("days", "生产天数", "text"),
    ("weight", "月产量(斤)", "weight"), ("price_a", "实际单价(元/斤)", "price"), ("price_s", "标准单价(元/斤)", "price"),
    ("diff_price", "差额(元/斤)", "price"), ("diff_amount", "月差额(元)", "money"),
    ("prev_price_a", "上月实际单价", "price"), ("price_change", "比上月", "yield"),
    ("yield_a", "熟制出率·实际", "yield"), ("yield_s", "熟制出率·标准", "yield"), ("yield_dev", "偏差(百分点)", "pp"),
    ("alert_days", "超限天数", "text"), ("no_cost_days", "未出成本天数", "text"),
    ("loss_days", "全损次数", "text"), ("loss_cost", "全损金额(元)", "money"),
]
MONTH_MATERIAL_COLS = [
    ("name", "原料", "text"), ("category", "类别", "text"), ("days", "处理天数", "text"),
    ("qty1", "1级领用(斤)", "weight"), ("qty2", "2级称重(斤)", "weight"),
    ("yield_a", "月出率·实际", "yield"), ("yield_s", "月出率·标准", "yield"), ("yield_dev", "偏差(百分点)", "pp"),
    ("impact", "出率偏差金额(元)", "money"), ("alert_days", "超限天数", "text"), ("std_days", "缺处理记录天数", "text"),
]


@router.get("/monthly")
def get_monthly(month: str, db: Session = Depends(get_db)):
    return monthly_report(db, month)


def _fill_sheet(ws, title: str, columns, rows, sub: str):
    """和 table_xlsx 一样的表格样式，写进指定的工作表（月度汇总一个文件两张表）。"""
    ws.append([title + (f"（{sub}）" if sub else "")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([c[1] for c in columns])
    for i in range(1, len(columns) + 1):
        c = ws.cell(2, i)
        c.fill, c.font = HEAD_FILL, Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = 14 if columns[i - 1][2] != "text" else 18
    for r in rows:
        ws.append([to_cell(r.get(k), t) for k, _, t in columns])
        for i, (k, _, t) in enumerate(columns, start=1):
            if t in FMT:
                ws.cell(ws.max_row, i).number_format = FMT[t]
        if r.get("alert_days"):
            for i in range(1, len(columns) + 1):
                ws.cell(ws.max_row, i).font = RED
    ws.freeze_panes = "A3"


@router.get("/monthly/export")
def export_monthly(month: str, db: Session = Depends(get_db)):
    rep = monthly_report(db, month)
    sub = f"{rep['start']} 至 {rep['end']}"
    wb = Workbook()
    ws = wb.active
    ws.title = "产品"
    _fill_sheet(ws, rep["title"] + " · 产品", rep["product_columns"], rep["products"], sub)
    sm = rep["summary"]
    ws.append([])
    for label, v, fmt in [("月总产量(斤)", sm["output"], "0.00"), ("生产天数", sm["production_days"], "0"),
                          ("成本差额合计(元)", sm["diff_amount"], "0.00"), ("原料出率偏差金额合计(元)", sm["material_impact"], "0.00"),
                          ("全损金额合计(元)", sm["loss_cost"], "0.00"), ("能耗费用合计(元)", sm["energy_cost"], "0.00"),
                          ("吨产品能耗费用(元/吨)", sm["energy_cost_per_ton"], "0.00")]:
        ws.append([label, to_cell(v, "money")])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).number_format = fmt
    _fill_sheet(wb.create_sheet("原料出率"), rep["title"] + " · 原料出率", rep["material_columns"], rep["materials"], sub)
    return respond(wb, rep["title"])
