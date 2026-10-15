"""月度汇总：月单价 = 月总成本 ÷ 月产量（不是日单价平均），月出率 = Σ产出 ÷ Σ投入，全损计入，导出两张表。"""
import io
from decimal import Decimal

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)


def D(v):
    return Decimal(v) if v is not None else None


def test_monthly_matches_daily(demo):
    rep = client.get("/api/reports/monthly", params={"month": "2026-10"}).json()
    daily = client.get("/api/reports/cost", params={"start": "2026-10-01", "end": "2026-10-31"}).json()["rows"]
    p = next(x for x in rep["products"] if x["name"] == "粉丝馅料")
    rows = [r for r in daily if r["name"] == "粉丝馅料"]
    assert p["days"] == len(rows) == 3
    w = sum(D(r["weight"]) for r in rows)
    assert D(p["weight"]) == w
    assert D(p["price_a"]) == sum(D(r["final_a"]) * D(r["weight"]) for r in rows) / w  # 按产量加权
    assert D(p["diff_amount"]) == sum(D(r["diff_amount"]) for r in rows)
    assert D(rep["summary"]["diff_amount"]) == sum(D(x["diff_amount"]) for x in rep["products"])

    m = next(x for x in rep["materials"] if x["name"] == "姜")
    assert D(m["yield_a"]) == D(m["qty2"]) / D(m["qty1"])
    assert m["days"] == 3 and D(m["impact"]) > 0  # 示例里姜的出率低于标准，多耗了原料


def test_material_impact_formula(demo):
    yr = client.get("/api/reports/yield", params={"start": "2026-10-12", "end": "2026-10-14", "all_materials": True}).json()
    rep = client.get("/api/reports/monthly", params={"month": "2026-10"}).json()
    m = next(x for x in rep["materials"] if x["name"] == "姜")
    # 按原料逐天：1 级单价 ×（1 级领用 − 2 级称重 ÷ 标准出率）
    from sqlalchemy import select
    from app.models import CalcMaterial, Material
    mid = demo.scalars(select(Material.id).where(Material.name == "姜")).one()
    rows = demo.scalars(select(CalcMaterial).where(CalcMaterial.material_id == mid,
                                                   CalcMaterial.yield_source == "actual")).all()
    expect = sum(r.price1 * (r.qty1 - r.qty2 / r.std_yield) for r in rows)
    assert D(m["impact"]) == expect
    assert yr["rows"]


def test_loss_and_prev_month(demo):
    d = client.get("/api/daily/2026-10-13").json()
    fs = next(r for r in d["production"] if r["name"] == "粉丝馅料")
    r = client.put("/api/daily/2026-10-13/production", json=[{**fs, "output_weight": "0"}])
    assert r.status_code == 200, r.text
    rep = client.get("/api/reports/monthly", params={"month": "2026-10"}).json()
    p = next(x for x in rep["products"] if x["name"] == "粉丝馅料")
    assert p["loss_days"] == 1 and D(p["loss_cost"]) > 0 and p["no_cost_days"] == 1
    assert D(rep["summary"]["loss_cost"]) == D(p["loss_cost"])
    nov = client.get("/api/reports/monthly", params={"month": "2026-11"}).json()
    assert nov["products"] == [] and nov["prev_month"] == "2026-10"


def test_monthly_export_and_bad_month(demo):
    r = client.get("/api/reports/monthly/export", params={"month": "2026-10"})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["产品", "原料出率"]
    assert wb["产品"]["A2"].value == "产品" and wb["原料出率"]["A2"].value == "原料"
    assert client.get("/api/reports/monthly", params={"month": "2026-13"}).status_code == 400
    assert client.get("/api/reports/monthly", params={"month": "abc"}).status_code == 400
