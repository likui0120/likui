"""接口冒烟测试：录入、导入、报表、导出都能跑通。"""
import io

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app

client = TestClient(app)
DAY = "2026-10-13"


def xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_daily_roundtrip_and_reports(demo):
    d = client.get(f"/api/daily/{DAY}").json()
    assert d["production"] and d["meters"] and d["packing"]
    fs = next(r for r in d["production"] if r["name"] == "粉丝馅料")
    r = client.put(f"/api/daily/{DAY}/production", json=[{**fs, "output_weight": "1100"}])
    assert r.status_code == 200, r.text
    assert any(i["kind"] == "alert" and i["object_name"] == "粉丝馅料" for i in r.json()["issues"])

    bad = client.put(f"/api/daily/{DAY}/production", json=[{**fs, "batches": "0", "output_weight": "1"}])
    assert bad.status_code == 400

    y = client.get("/api/reports/yield", params={"start": DAY, "end": DAY}).json()
    assert any(row["alert"] for row in y["rows"])
    c = client.get("/api/reports/cost", params={"start": DAY, "end": DAY}).json()
    assert len(c["rows"]) == 5
    pid = fs["product_id"]
    t = client.get("/api/reports/trace", params={"product_id": pid, "day": DAY}).json()
    assert len(t["trace"]["lines"]) == 13 and t["attribution"]

    iss = client.get("/api/reports/issues", params={"start": DAY, "end": DAY}).json()
    alert = next(x for x in iss["rows"] if x["kind"] == "alert")
    client.put("/api/reports/notes", json={"date": DAY, "item_key": alert["item_key"], "note": "测试原因"})
    iss = client.get("/api/reports/issues", params={"start": DAY, "end": DAY}).json()
    assert any(x["note"] == "测试原因" for x in iss["rows"])

    ed = client.get("/api/reports/energy-daily", params={"start": "2026-10-12", "end": "2026-10-14"}).json()
    assert len(ed["rows"]) == 3 and ed["rows"][0]["gas_t"]
    em = client.get("/api/reports/energy-month", params={"month": "2026-10"}).json()
    assert len(em["days"]) == 31
    tr = client.get("/api/reports/trend", params={"kind": "product_cost", "ref_id": pid, "stage": "final",
                                                 "start": "2026-10-12", "end": "2026-10-14"}).json()
    assert len(tr["points"]) == 3

    for name, params in [("yield", {"start": DAY, "end": DAY}), ("cost", {"start": DAY, "end": DAY}),
                         ("issues", {"start": DAY, "end": DAY}), ("energy-daily", {"start": DAY, "end": DAY}),
                         ("energy-month", {"month": "2026-10"}), ("trace", {"product_id": pid, "day": DAY})]:
        r = client.get(f"/api/reports/export/{name}", params=params)
        assert r.status_code == 200, (name, r.text)
        load_workbook(io.BytesIO(r.content))

    dash = client.get("/api/dashboard", params={"day": DAY}).json()
    assert dash["products"] == 5


def test_meter_validation(demo):
    d = client.get("/api/daily/2026-10-14").json()
    m = next(x for x in d["meters"] if x["mode"] == "reading" and x["plant"] == "+")
    r = client.put("/api/daily/2026-10-14/meters", json=[{**m, "reading": "1"}])
    assert any("小于上期" in i["message"] for i in r.json()["issues"])


def test_import_tplus(demo):
    f = xlsx([["T+ 存货单价表"], ["存货编码", "存货名称", "计量单位", "数量", "单价", "金额"],
              ["", "姜", "公斤", 10, 9, 90], ["", "姜", "公斤", 30, 10, 300], ["", "不存在的料", "斤", 1, 1, 1],
              ["", "合计", "", "", "", 391]])
    r = client.post("/api/import/tplus", data={"day": "2026-10-15", "unit": "jin"},
                    files={"file": ("tplus.xlsx", f, "application/octet-stream")})
    assert r.status_code == 200, r.text
    body = r.json()
    assert [u["name"] for u in body["unmatched"]] == ["不存在的料"]
    ginger = next(m for m in body["matched"] if m["name"] == "姜")
    # 加权：390 元 ÷ 80 斤
    assert abs(float(ginger["price"]) - 390 / 80) < 1e-9


def test_import_materials_products_boms(db):
    from app.seed import init_base
    init_base(db)
    db.commit()
    f = xlsx([["编码", "名称", "类别", "是否需要处理", "标准出率"], ["M01", "香菜", "蔬菜", "是", "80%"],
              ["", "姜", "蔬菜", "是", "0.86"]])
    r = client.post("/api/import/materials", data={"effective_date": "2026-01-01"},
                    files={"file": ("m.xlsx", f, "application/octet-stream")}).json()
    assert r["created"] == 1 and r["updated"] == 1
    f = xlsx([["名称", "类型", "班组", "是否制冷", "是否包装", "熟制标准出率"], ["香菜馅", "成品馅", "拌馅组", "否", "是", "99%"]])
    r = client.post("/api/import/products", data={"effective_date": "2026-01-01"},
                    files={"file": ("p.xlsx", f, "application/octet-stream")}).json()
    assert r["created"] == 1
    f = xlsx([["产品", "行号", "投料", "数量"], ["香菜馅", 1, "香菜", 10], [None, 2, "姜", 5],
              ["粉丝馅料", 1, "不存在", 3]])
    r = client.post("/api/import/boms", data={"effective_date": "2026-01-01"},
                    files={"file": ("b.xlsx", f, "application/octet-stream")}).json()
    assert r["imported"] == [{"product": "香菜馅", "lines": 2, "total": "15"}]
    assert r["unmatched"][0]["input"] == "不存在"
    assert client.get("/api/import/template/boms").status_code == 200


def test_master_crud_triggers_recalc(demo):
    prods = client.get("/api/products").json()
    pid = next(p["id"] for p in prods if p["name"] == "粉丝馅料")
    before = client.get("/api/reports/cost", params={"start": DAY, "end": DAY}).json()["rows"]
    stds = [s for s in client.get("/api/product-std-yields").json() if s["product_id"] == pid]
    r = client.put(f"/api/product-std-yields/{stds[0]['id']}", json={"cook_yield": "0.95"})
    assert r.status_code == 200
    after = client.get("/api/reports/cost", params={"start": DAY, "end": DAY}).json()["rows"]
    fb = next(x for x in before if x["product_id"] == pid)
    fa = next(x for x in after if x["product_id"] == pid)
    assert fb["final_s"] != fa["final_s"] and fb["final_a"] == fa["final_a"]
    # 被引用的原料不能删
    mats = client.get("/api/materials").json()
    ginger = next(m for m in mats if m["name"] == "姜")
    assert client.delete(f"/api/materials/{ginger['id']}").status_code == 400
